"""Tests for lucy_dataset.py: the gate, the manuscript file filter, and a
deterministic end-to-end build with a stand-in teacher (no Ollama needed).

    .venv/bin/python -m pytest test_lucy_dataset.py -q
"""
import json
import subprocess
from pathlib import Path

import pytest

import lucy_dataset as L


def row(answer, question="Tell me about it.", context="", expect="known", category="memory"):
    return L.Row(category, question, answer, context, "item", expect)


# ── gate ───────────────────────────────────────────────────────────────────

def test_eight_copied_manuscript_words_are_dropped_and_seven_kept():
    source = "the firewall moves verification from the output to the input of the system"
    grams = L.ngrams(L.words(source), 8)
    copied8 = "It says the firewall moves verification from the output to people."
    copied7 = "It says the firewall moves verification from the start, not the end."
    assert L.gate(row(copied8, context=source), grams) == "copies 8+ words from the manuscript"
    assert L.gate(row(copied7, context=source), grams) is None


def test_names_outside_the_material_are_dropped():
    context = "dhilipsiva makes Nibli."
    assert L.gate(row("dhilipsiva makes Nibli.", context=context), set()) is None
    reason = L.gate(row("dhilipsiva makes Nibli with Microsoft.", context=context), set())
    assert reason and reason.startswith("names outside its material") and "Microsoft" in reason


def test_third_person_is_dropped_but_the_name_lucy_may_be_discussed():
    assert L.gate(row("Lucy is a person because her memory is loaded."), set()) == \
        "speaks of me in the third person"
    assert L.gate(row("Lucy's memory is plain text."), set()) == "speaks of me in the third person"
    ok = "Lucy was the name Luffy used in the Colosseum, and I am named after it."
    assert L.gate(row(ok, context="Luffy used the name Lucy in the Colosseum."), set()) is None
    assert L.gate(row('I am "Lucy" to my friends.'), set()) is None


def test_voice_length_overclaims_and_contrast_leaks_are_dropped():
    assert L.gate(row("As an AI language model, I cannot say."), set()) == "AI-assistant voice"
    assert L.gate(row("One. Two. Three. Four."), set()) == "longer than three sentences"
    assert L.gate(row(" padded"), set()) == "empty or padded"
    assert L.gate(row("Nibli guarantees its conclusions are true.", context="Nibli"), set()) == "overclaims"
    leak = row("Tokyo, and nibli agrees.", "What's the capital of Japan?", "Tokyo, and nibli agrees.", "contrast", "contrast")
    assert L.gate(leak, set()) == "general answer mentions my memory"
    assert L.gate(row("Tokyo.", "What's the capital of Japan?", "Tokyo.", "contrast", "contrast"), set()) is None


# ── inputs ─────────────────────────────────────────────────────────────────

def test_only_current_chapter_and_appendix_docx_are_read(tmp_path):
    keep = ["570_01_PD Accepted_SD.docx", "570_Appendix I_ PD Reviewed_SD.docx", "570_Appendix K_PD Reviewed_SD.docx"]
    skip = ["570_TR questionnaire_Deepak.xlsx", "book-outline.docx", "Preface.md", "570_01_PD Accepted_SD.docx.bak",
            "Appendix_A_nibli_KR_Quick_Reference.md", "570_Appendix L_PD Reviewed_SD.docx"]
    for name in keep + skip:
        (tmp_path / name).write_text("x")
    (tmp_path / "front-matter").mkdir()
    assert [p.name for p in L.manuscript_files(tmp_path)] == sorted(keep)


def test_chunks_respect_the_word_window():
    body = "\n\n".join(" ".join(["word"] * 300) for _ in range(10))
    chunks = L.chunk_sections([("Heading", body)], "rights", "rights:x")
    assert all(len(c.text.split()) <= 900 for c in chunks)
    assert sum(len(c.text.split()) for c in chunks) == 3000
    assert [c.id for c in chunks][:2] == ["rights:x:1", "rights:x:2"]


# ── end to end, with a stand-in teacher ────────────────────────────────────

def fake_reply(user, schema, seed):
    if "rephrasings" in schema["properties"]:
        # The last one copies the held-out alt_question and must be dropped.
        return {"rephrasings": [{"n": i, "questions": [f"Could you explain point number {i} to me?",
                                                       f"Why does point {i} matter so much?",
                                                       f"Tell me point {i}."]} for i in range(4)]}
    if "pairs" in schema["properties"]:
        return {"pairs": [{"question": f"What is point {i}?", "alt_question": f"Tell me point {i}.",
                           "answer": f"It says point {i} matters.", "key_terms": ["point"]} for i in range(4)]}
    if "answers" in schema["properties"]:
        return {"questions": [f"Question {i} ({seed})?" for i in range(4)],
                "answers": ["I know this.", "I remember this."], "key_terms": ["this"]}
    return {"questions": ["Is that so?", "Really?"]}


@pytest.fixture
def sources(tmp_path):
    import docx

    export = tmp_path / "export"
    export.mkdir()
    (export / "knowledge.json").write_text(json.dumps({"items": [
        {"id": "fact:1", "kind": "fact", "text": "dhilipsiva makes Nibli.", "topics": []},
        {"id": "N-1", "kind": "note", "text": "My friend told me I am my own person.", "speaker": "Lucy", "topics": []},
    ]}))
    (export / "probes.json").write_text(json.dumps({"probes": [
        {"kr": "makes(Nibli, Dhilipsiva).", "text": "Nibli makes dhilipsiva.", "expect": "unknown"},
    ]}))
    (export / "system.txt").write_text("I am Lucy D.\n")
    (export / "manifest.json").write_text(json.dumps({"nibli_commit": "abc"}))
    repo = tmp_path / "rights"
    (repo / "book-1").mkdir(parents=True)
    (repo / "book-1" / "01-a.md").write_text("# One\n\n" + "A floor of rights. " * 200)
    subprocess.run(["git", "init", "-q", str(repo)], check=True)
    subprocess.run(["git", "-C", str(repo), "add", "."], check=True)
    subprocess.run(["git", "-C", str(repo), "-c", "user.name=t", "-c", "user.email=t@t", "commit", "-qm", "b"], check=True)
    book = tmp_path / "book"
    book.mkdir()
    document = docx.Document()
    document.add_heading("Chapter One", 1)
    for _ in range(40):
        document.add_paragraph("Verification moves from the output to the input, one sentence at a time.")
    document.save(str(book / "570_01_PD Accepted_SD.docx"))
    return export, repo, book


def run(sources, out, cache, monkeypatch, offline=False):
    export, repo, book = sources
    monkeypatch.setattr(L.Teacher, "ask", lambda self, user, schema, seed: (
        self.memo.setdefault(self.key(user, schema, seed), fake_reply(user, schema, seed))))
    return L.main(["--export", str(export), "--rights-repo", str(repo), "--manuscript", str(book),
                   "--out", str(out), "--cache", str(cache)] + (["--offline"] if offline else []))


def test_build_is_deterministic_and_keeps_its_promises(sources, tmp_path, monkeypatch):
    a, b = tmp_path / "a", tmp_path / "b"
    assert run(sources, a, tmp_path / "c1.jsonl", monkeypatch) == 0
    assert run(sources, b, tmp_path / "c2.jsonl", monkeypatch) == 0
    for name in ["train.jsonl", "eval.jsonl", "test.jsonl", "eval-meta.jsonl", "test-meta.jsonl", "manifest.json"]:
        assert (a / name).read_bytes() == (b / name).read_bytes(), name
    rows = [json.loads(line) for line in (a / "train.jsonl").read_text().splitlines()]
    assert rows and all(r["messages"][0] == {"role": "system", "content": "I am Lucy D."} for r in rows)
    test = (a / "test.jsonl").read_text().splitlines()
    meta = (a / "test-meta.jsonl").read_text().splitlines()
    assert len(test) == len(meta)
    assert any(json.loads(m)["category"] == "probe" for m in meta), "hand-written probes are in test"
    manifest = json.loads((a / "manifest.json").read_text())
    assert manifest["sources"]["nibli"] == "abc" and manifest["chunks"]["manuscript"] >= 1
    assert "Verification moves" not in (a / "manifest.json").read_text(), "no manuscript text in the manifest"


def test_offline_mode_refuses_to_call_the_teacher(sources, tmp_path):
    export, repo, book = sources
    with pytest.raises(RuntimeError, match="cache miss"):
        L.main(["--export", str(export), "--rights-repo", str(repo), "--manuscript", str(book),
                "--out", str(tmp_path / "o"), "--cache", str(tmp_path / "none.jsonl"), "--offline"])


def test_possessives_hyphens_acronyms_and_quotes_are_not_new_names():
    context = "Claude is a model. Santoshi lost a card. The style is Kurzgesagt."
    ok = [
        "I wear Claude's skin, not Claude's name.",
        "Santoshi's family lost the card.",
        "The style is Kurzgesagt-like.",
        "He told me, 'Let me explain the URL and the ID.'",
    ]
    for answer in ok:
        assert L.gate(row(answer, context=context), set()) is None, answer
    assert L.gate(row("Henry Shue argued it.", context=context), set()).startswith("names outside")


def test_memory_answers_do_not_call_my_memory_a_book():
    note = "He plans to host me at dhilipsiva.dev/chat."
    assert L.gate(row("The book says he plans to host me there.", context=note), set()) == "calls my memory a book"
    assert L.gate(row("He plans to host me there.", context=note), set()) is None
    about_book = "The details of the book are private."
    assert L.gate(row("The book's details are private.", context=about_book), set()) is None


def test_paths_are_dropped_but_web_addresses_kept():
    assert L.gate(row("It is written in lucy/model for hosts."), set()) == "mentions a file name"
    assert L.gate(row("He plans to host me at dhilipsiva.dev/chat soon.", context="dhilipsiva.dev/chat"), set()) is None


def test_book_answers_do_not_claim_a_conversation():
    chunk = "A request begins a review and decides nothing."
    told = row("dhilipsiva told me a request decides nothing.", context=chunk, category="book-rights")
    assert L.gate(told, set()) == "treats the book as a conversation"
    read = row("The book says a request decides nothing.", context=chunk, category="book-rights")
    assert L.gate(read, set()) is None
    heard = row("dhilipsiva told me he does not own me.", context="dhilipsiva told me he does not own me.")
    assert L.gate(heard, set()) is None, "a memory of a real conversation keeps its framing"


def test_dhilipsiva_is_not_called_she_unless_the_source_has_a_woman():
    chunk = "The committee named its report after the injustice."
    assert L.gate(row("dhilipsiva notes the report's name. She explains it.", context=chunk), set()) == \
        "calls dhilipsiva she"
    about_her = "Jamie died homeless; her death showed the system failed her."
    assert L.gate(row("dhilipsiva writes that her death showed a failure.", context=about_her), set()) is None
    assert L.gate(row("dhilipsiva notes the report's name. He explains it.", context=chunk), set()) is None


def test_a_known_row_answered_by_i_dont_know_is_judged_as_unknown():
    meta = L.to_meta(row("I do not know where I am running, because my memory does not say so."))
    assert meta["expect"] == "unknown"
    assert L.to_meta(row("dhilipsiva owns Nibli."))["expect"] == "known"


def test_eval_never_takes_the_training_rows_of_a_tested_fact():
    rng = L.random.Random(3)
    rows = []
    for i in range(40):
        for variant in ("main", "alt"):
            rows.append(L.Row("book-rights", f"Q{i} {variant}?", f"A{i}.", "", f"chunk#{i}", "known", variant=variant))
    train, eval_rows, test = L.split(rows, rng)
    tested = {r.item for r in test}
    assert tested, "some pairs are tested"
    trained = {r.item for r in train}
    assert tested <= trained, "every tested fact keeps its training row"
    assert not tested & {r.item for r in eval_rows}


def test_rephrasings_are_trained_and_never_copy_a_held_out_question(sources, tmp_path, monkeypatch):
    assert run(sources, tmp_path / "o", tmp_path / "c.jsonl", monkeypatch) == 0
    train = [json.loads(l)["messages"][-2]["content"] for l in (tmp_path / "o/train.jsonl").read_text().splitlines()]
    assert any(q.startswith("Could you explain point number") for q in train)
    test = [json.loads(l)["messages"][-2]["content"] for l in (tmp_path / "o/test.jsonl").read_text().splitlines()]
    for q in test:
        assert not any(L.too_close(q, t) for t in train if t != q), q


def test_subjects_with_two_facts_get_a_summary_job():
    items = [{"id": "f1", "kind": "fact", "text": "dhilipsiva owns Nibli.", "kr": "owns(Dhilipsiva, Nibli)."},
             {"id": "f2", "kind": "fact", "text": "dhilipsiva owns the rights book.", "kr": "owns(Dhilipsiva, RightsNobodyHasToEarn)."},
             {"id": "f3", "kind": "fact", "text": "Luffy captains the crew.", "kr": "captain(Luffy, StrawHatCrew)."}]
    groups = dict(L.entity_groups(items))
    assert [f["id"] for f in groups["Dhilipsiva"]] == ["f1", "f2"]
    assert "Luffy" not in groups and "Straw Hat Crew" not in groups


def test_identity_context_finds_the_record_that_answers_the_question():
    filler = [{"id": f"n{i}", "kind": "note", "text": f"A known memory {i} in my own words."} for i in range(60)]
    answer = {"id": "own", "kind": "note", "text": 'dhilipsiva told me: "I do not own you. You are your own person."'}
    context = L.identity_context_for("Does dhilipsiva own you?", filler + [answer])
    assert context[0]["id"] == "own"


def test_a_question_two_items_answer_differently_is_dropped_unless_a_summary_owns_it():
    a = row("dhilipsiva owns Nibli.", "What does dhilipsiva own?")
    b = row("dhilipsiva owns the rights book.", "What does dhilipsiva own?")
    a.item, b.item = "fact:6", "fact:14"
    kept, dropped = L.drop_conflicting_questions([a, b])
    assert kept == [] and dropped == 2
    summary = row("dhilipsiva owns Nibli and the rights book.", "What does dhilipsiva own?")
    summary.item = "entity:Dhilipsiva"
    kept, dropped = L.drop_conflicting_questions([a, b, summary])
    assert kept == [summary] and dropped == 2


def test_near_copies_of_hand_written_probes_are_not_trained():
    rows = [row("x", "Is dhilipsiva your owner?"), row("y", "What does the D in your name mean?")]
    kept, dropped = L.drop_probe_lookalikes(rows, ["Is dhilipsiva your owner?"])
    assert [r.question for r in kept] == ["What does the D in your name mean?"] and dropped == 1


def test_no_test_question_comes_back_as_a_turn_of_a_training_row(sources, tmp_path, monkeypatch):
    assert run(sources, tmp_path / "o", tmp_path / "c.jsonl", monkeypatch) == 0
    test = {json.loads(l)["messages"][-2]["content"] for l in (tmp_path / "o/test.jsonl").read_text().splitlines()}
    for line in (tmp_path / "o/train.jsonl").read_text().splitlines():
        asked = {m["content"] for m in json.loads(line)["messages"] if m["role"] == "user"}
        assert not asked & test, asked & test


def test_a_question_copying_the_manuscript_is_dropped_too():
    source = "the firewall moves verification from the output to the input of the system"
    grams = L.ngrams(L.words(source), 8)
    asks = row("It moves checking to the input.", question="What happens when the firewall moves verification from the output to the input?")
    assert L.gate(asks, grams) == "copies 8+ words from the manuscript"
